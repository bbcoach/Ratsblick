"""Erzeugt data/gebiete-rlp.json aus dem Gemeindeverzeichnis von Destatis (GV-ISys, Auszug).

Aufruf:
  pip install openpyxl
  curl -L -o gv.xlsx "https://www.destatis.de/DE/Themen/Laender-Regionen/Regionales/Gemeindeverzeichnis/Administrativ/Archiv/GVAuszugQ/AuszugGV2QAktuell.xlsx?__blob=publicationFile"
  python3 scripts/gemeindeverzeichnis.py gv.xlsx

Schlüssel: Kreis = 5 Stellen (07131), Verbandsgemeinde = 9 Stellen (071315001), Gemeinde = AGS mit 8 Stellen (07131007).
"""
import json
import sys

import openpyxl

LAND = '07'


def kreisname(name: str, kreisfrei: bool) -> str:
    if kreisfrei:
        return name.replace(', kreisfreie Stadt', '')
    return name if name.lower().endswith('kreis') or name.startswith('Eifelkreis') else f'Landkreis {name}'


def main(path: str) -> None:
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[1]
    stand = ws.cell(row=1, column=1).value
    kreise, vgs, gemeinden = {}, {}, []
    for row in ws.iter_rows(min_row=8, values_only=True):
        satz, tk, land, rb, kreis, vb, gem, name = row[:8]
        if land != LAND or kreis is None:
            continue
        k = f'{land}{rb}{kreis}'
        if satz == '40':
            kreise[k] = {'id': k, 'name': kreisname(name, tk == '41'), 'art': 'Kreisfreie Stadt' if tk == '41' else 'Landkreis'}
        elif satz == '50' and tk == '53':
            vgs[f'{k}{vb}'] = {'id': f'{k}{vb}', 'name': f'Verbandsgemeinde {name}', 'kreis': k}
        elif satz == '60' and tk in ('61', '63', '64'):
            stadt = name.endswith(', Stadt')
            rein = name.replace(', Stadt', '')
            vg = f'{k}{vb}' if vb.startswith('5') else None
            if tk == '61':
                art = 'Kreisfreie Stadt'
            elif stadt:
                art = 'Stadt'
            else:
                art = 'Ortsgemeinde' if vg else 'Gemeinde'
            gemeinden.append({
                'id': f'{k}{gem}', 'name': rein, 'art': art,
                'plz': row[13] or None, 'einwohner': row[9] or None,
                'kreis': None if tk == '61' else k, 'vg': vg,
            })
    out = {
        'quelle': 'Statistisches Bundesamt, Gemeindeverzeichnis-Informationssystem (GV-ISys), Datenlizenz Deutschland 2.0',
        'stand': stand,
        'kreise': sorted(kreise.values(), key=lambda x: x['id']),
        'verbandsgemeinden': sorted(vgs.values(), key=lambda x: x['id']),
        'gemeinden': sorted(gemeinden, key=lambda x: x['id']),
    }
    with open('data/gebiete-rlp.json', 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=0, separators=(',', ':'))
    print(f"{len(kreise)} Kreise/kreisfreie Städte, {len(vgs)} Verbandsgemeinden, {len(gemeinden)} Gemeinden")


if __name__ == '__main__':
    main(sys.argv[1])
